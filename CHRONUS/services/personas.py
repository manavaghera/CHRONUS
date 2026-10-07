"""
Personas — the people CHRONUS can simulate.

* Pretrained personas ship with the repo: models/<id>/persona.json (Elon Musk).
* Custom personas are built by users through the "Create your model" flow
  (consent -> upload documents -> structured interview -> build) and live in
  personas/<id>/ — git-ignored, because they hold private personal data.

Every persona has its own ChromaDB collection, so one person's memories can
never surface in another person's answers, and deleting a custom persona
removes its vectors, uploads and record in one step (the paper's deletion
requirement).
"""

from __future__ import annotations

import contextvars
import hashlib
import json
import os
import re
import secrets
import shutil
import tempfile
import threading
import time
from datetime import datetime
from pathlib import Path

from services import chroma_index
from services.timeline import year_of

ROOT = Path(__file__).resolve().parent.parent
PRETRAINED_DIR = ROOT / "models"
CUSTOM_DIR = ROOT / "personas"

PERSONA_ID_PATTERN = r"^[a-z0-9_]{3,48}$"
UPLOAD_TYPES = {".txt", ".md", ".pdf", ".csv", ".json", ".docx"}
# Voice notes and recordings: transcribed on this computer (services/stt.py,
# needs faster-whisper) and stored as their spoken words
AUDIO_TYPES = {".wav", ".mp3", ".m4a", ".ogg", ".webm", ".flac", ".aac"}
# A new memory this close (cosine distance) to one from another file is the
# same text uploaded twice (a letter in two documents): skipped
DUPLICATE_DISTANCE = 0.03
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
# Data sufficiency check before a custom model can be built (paper, Input
# Data Layer): with fewer memories it can barely answer anything.
MIN_MEMORIES_TO_BUILD = 10
EMBED_BATCH = 64

DEFAULT_STYLE_NOTES = (
    "Talk like a person in a conversation, not like an essay: plain words and "
    "concrete details instead of big abstract phrases."
)

# persona.json files are read-modify-written from concurrent requests
_lock = threading.Lock()


def _now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _folder(persona: dict) -> Path:
    base = PRETRAINED_DIR if persona["kind"] == "pretrained" else CUSTOM_DIR
    return base / persona["id"]


# Signed-in user for this request when the server has accounts
# (CHRONUS_USERS, services/access.py); None = single-user server.
current_user: contextvars.ContextVar[str | None] = contextvars.ContextVar("chronus_user", default=None)


def _visible(persona: dict) -> bool:
    """Custom models belong to the account that made them. Pretrained ones,
    and models made before accounts existed (no owner), are shared."""
    user = current_user.get()
    return user is None or persona.get("kind") != "custom" or persona.get("owner") in (None, user)


def _read_json(path: Path) -> dict:
    """persona.json is replaced atomically (save_persona), but on Windows the
    swap can briefly deny reads; retry instead of failing the request."""
    for attempt in range(20):
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except (PermissionError, json.JSONDecodeError):
            if attempt == 19:
                raise
            time.sleep(0.01 * (attempt + 1))
    raise AssertionError("unreachable")


def load_persona(persona_id: str, any_owner: bool = False) -> dict | None:
    if not re.match(PERSONA_ID_PATTERN, persona_id or ""):
        return None
    for base in (PRETRAINED_DIR, CUSTOM_DIR):
        path = base / persona_id / "persona.json"
        if path.exists():
            persona = _read_json(path)
            return persona if any_owner or _visible(persona) else None
    return None


def list_personas(any_owner: bool = False) -> list[dict]:
    found = []
    for base in (PRETRAINED_DIR, CUSTOM_DIR):
        for path in sorted(base.glob("*/persona.json")) if base.exists() else []:
            persona = _read_json(path)
            if any_owner or _visible(persona):
                found.append(persona)
    return found


def atomic_write_text(path: Path, text: str) -> None:
    """Write via a temp file in the same folder, then swap it in, so readers
    see the old file or the new one, never a half-written one (8 of 25
    interview answers once failed with 500s while persona.json was rewritten
    in place). Windows refuses the swap while a reader has the file open, so
    retry briefly."""
    fd, tmp = tempfile.mkstemp(dir=path.parent, prefix=f".{path.name}.", suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            f.write(text)
        for attempt in range(50):
            try:
                os.replace(tmp, path)
                return
            except PermissionError:
                if attempt == 49:
                    raise
                time.sleep(0.01 * (attempt + 1))
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)


def save_persona(persona: dict) -> None:
    folder = _folder(persona)
    folder.mkdir(parents=True, exist_ok=True)
    atomic_write_text(folder / "persona.json", json.dumps(persona, indent=2, ensure_ascii=False))


def create_custom_persona(
    name: str, description: str, relationship: str, allow_cloud_llm: bool, consent_statement: str,
    consent_language: str = "en",
    memorial: bool = False,
) -> dict:
    slug = re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")[:30] or "persona"
    persona_id = f"{slug}_{secrets.token_hex(3)}"
    persona = {
        "id": persona_id,
        "name": name.strip(),
        "kind": "custom",
        "status": "draft",
        "description": description.strip(),
        "collection": f"persona_{persona_id}",
        # Off by default: AI voice sends evidence excerpts to the cloud LLM
        # provider, and the paper's privacy model is local-first.
        "allow_cloud_llm": allow_cloud_llm,
        "style_notes": DEFAULT_STYLE_NOTES,
        "relationship": relationship,  # creator's relationship to the person
        "memorial": memorial,  # the person has died (services/wellbeing.py)
        "owner": current_user.get(),  # account that made it (None on a single-user server)
        # the exact words agreed to, in the language they were read in (services/consent_text.py)
        "consent": {"statement": consent_statement, "language": consent_language, "given_at": _now()},
        "created_at": _now(),
        "uploads": [],
        "interview_answered": [],
    }
    with _lock:
        save_persona(persona)
    return persona


def get_collection(client, persona: dict):
    return client.get_or_create_collection(name=persona["collection"], metadata=chroma_index.COLLECTION_METADATA)


def _safe_filename(name: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._ -]", "_", Path(name).name).strip(" .")[:80]
    return cleaned or "upload.txt"


def _sanitize(value):
    """Chroma metadata only accepts str/int/float/bool."""
    if value is None:
        return ""
    return value if isinstance(value, (str, int, float, bool)) else str(value)


def _parse_docx(path: Path, merge_sources) -> list[dict]:
    """Word document -> records, via the same paragraph chunking as .txt/.md."""
    import docx  # python-docx

    try:
        document = docx.Document(str(path))
    except Exception as e:  # not a real .docx (zip) file
        raise ValueError(f"Couldn't read this Word file: {e}")
    blocks = [p.text for p in document.paragraphs]
    for table in document.tables:
        blocks += [" | ".join(cell.text for cell in row.cells) for row in table.rows]
    text_path = path.with_name(path.name + ".txt")
    try:
        text_path.write_text("\n\n".join(b.strip() for b in blocks if b.strip()), encoding="utf-8")
        return merge_sources.parse_markdown(text_path)
    finally:
        text_path.unlink(missing_ok=True)


def ingest_document(persona: dict, collection, embedder, filename: str, data: bytes, authored_by: str,
                    progress=None) -> dict:
    """Parse, chunk and embed one uploaded document into the persona's memory.

    Re-uploading a file with the same name replaces its earlier memories.
    Memories identical to ones already stored from other files are skipped.
    *progress(fraction, message)*: called as the work advances (background
    uploads show it as a progress bar).
    Raises ValueError for unsupported, oversized or empty files.
    """
    import merge_sources  # parsers + chunker; deferred so importing this module stays cheap

    report = progress or (lambda fraction, message: None)
    name = _safe_filename(filename)
    ext = Path(name).suffix.lower()
    from services import originals

    accepted = UPLOAD_TYPES | AUDIO_TYPES | originals.IMAGE_TYPES
    if ext not in accepted:
        raise ValueError(f"Unsupported file type '{ext or name}'. Use: {', '.join(sorted(accepted))}")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError(f"File is larger than {MAX_UPLOAD_BYTES // (1024 * 1024)} MB")
    check_upload_content(ext, data)

    uploads = CUSTOM_DIR / persona["id"] / "uploads"
    uploads.mkdir(parents=True, exist_ok=True)
    path = uploads / name
    path.write_bytes(data)

    report(0.05, "Reading the file")
    try:
        if ext in AUDIO_TYPES:
            report(0.1, "Transcribing the recording")
            records = [{**r, "source_file": name} for r in originals.transcribe(data)]
        elif ext in originals.IMAGE_TYPES:  # a photo of a letter, read on this computer
            report(0.1, "Reading the photo")
            records = originals.read_image(data)
        elif ext in (".txt", ".md"):
            records = merge_sources.parse_markdown(path)
        elif ext == ".docx":
            records = _parse_docx(path, merge_sources)
        elif ext == ".csv":
            records = merge_sources.parse_csv_file(path)
        elif ext == ".json":
            records = merge_sources.parse_json_file(path)
        else:
            records = merge_sources.parse_pdf_file(str(path))
            if sum(len(r["text"]) for r in records) < 50:  # a scan with no text layer
                report(0.1, "Reading the scanned pages")
                records = originals.read_scanned_pdf(path)
    except ValueError:
        path.unlink(missing_ok=True)  # don't keep a file we couldn't use
        raise
    except Exception as e:  # a parser crashed on a damaged file: same answer, no stray file
        path.unlink(missing_ok=True)
        raise ValueError(f"Couldn't read this file ({type(e).__name__})")

    # Their own writing (letters, journals, emails) and their own recorded
    # voice are first-person evidence; anything about them is third-party
    # (services/provenance.py).
    if ext in AUDIO_TYPES:
        source_type = "voice_note" if authored_by == "self" else "written_about"
    else:
        source_type = "personal_writing" if authored_by == "self" else "written_about"
    units = []
    for record in records:
        record.update(source_type=source_type, source_name=Path(name).stem, source_file=name)
        units.extend(merge_sources.chunk_record(record))
    units = [u for u in units if len(u["text"].strip()) >= 20]
    # The same paragraph twice in one file (signatures, repeated headers)
    seen: set[str] = set()
    units = [u for u in units if not (" ".join(u["text"].lower().split()) in seen or seen.add(" ".join(u["text"].lower().split())))]
    if not units:
        path.unlink(missing_ok=True)
        raise ValueError("No readable text found in this file")

    collection.delete(where={"source_file": name})  # replace an earlier upload of this file
    ids, docs, metas = [], [], []
    for i, unit in enumerate(units):
        memory_id = "cu_" + hashlib.md5(f"{name}|{i}|{unit['text']}".encode("utf-8")).hexdigest()[:16]
        meta = {k: _sanitize(v) for k, v in unit.items() if k not in ("text", "char_count", "word_count")}
        meta.update(memory_id=memory_id, person=persona["id"], authored_by=authored_by,
                    importance_score=merge_sources.estimate_importance(unit))
        year = year_of(meta.get("date"))
        if year is not None:
            meta["year"] = year  # for time travel (services/timeline.py)
        ids.append(memory_id)
        docs.append(unit["text"])
        metas.append(meta)

    stored = duplicates = 0
    has_memories = collection.count() > 0
    for start in range(0, len(docs), EMBED_BATCH):
        batch = slice(start, start + EMBED_BATCH)
        vectors = embedder.encode(docs[batch], normalize_embeddings=True).tolist()
        keep = list(range(len(vectors)))
        if has_memories:
            near = collection.query(query_embeddings=vectors, n_results=1, include=["distances"])
            keep = [k for k, dists in enumerate(near["distances"]) if not (dists and dists[0] <= DUPLICATE_DISTANCE)]
        duplicates += len(vectors) - len(keep)
        if keep:
            collection.upsert(ids=[ids[batch][k] for k in keep], documents=[docs[batch][k] for k in keep],
                              metadatas=[metas[batch][k] for k in keep], embeddings=[vectors[k] for k in keep])
            stored += len(keep)
        report(0.15 + 0.8 * min(1.0, (start + EMBED_BATCH) / len(docs)), f"Embedded {min(start + EMBED_BATCH, len(docs))} of {len(docs)} passages")

    # All duplicates (the same letter saved twice) still counts as uploaded:
    # the file is theirs; it just added nothing new.
    entry = {"filename": name, "authored_by": authored_by, "memories": stored, "duplicates_skipped": duplicates,
             "kind": "audio" if ext in AUDIO_TYPES else "image" if ext in originals.IMAGE_TYPES else "document",
             "uploaded_at": _now()}
    with _lock:
        fresh = load_persona(persona["id"], any_owner=True) or persona
        fresh["uploads"] = [u for u in fresh.get("uploads", []) if u["filename"] != name] + [entry]
        save_persona(fresh)
    report(1.0, "Done")
    return entry


_MAGIC = {".pdf": (b"%PDF",), ".docx": (b"PK\x03\x04",), ".wav": (b"RIFF",), ".flac": (b"fLaC",), ".ogg": (b"OggS",),
          ".webm": (b"\x1aE\xdf\xa3",), ".mp3": (b"ID3", b"\xff\xfb", b"\xff\xf3", b"\xff\xf2")}


def check_upload_content(ext: str, data: bytes) -> None:
    """The file's bytes must match its name: a renamed executable or HTML
    page is refused before any parser sees it."""
    from services.originals import IMAGE_MAGIC

    head = data[:16]
    magic = {**_MAGIC, **IMAGE_MAGIC}
    if ext in magic and not head.startswith(magic[ext]):
        raise ValueError(f"This doesn't look like a real {ext} file")
    if ext == ".webp" and data[8:12] != b"WEBP":
        raise ValueError("This doesn't look like a real .webp file")
    if ext in (".txt", ".md", ".csv", ".json"):
        if b"\x00" in data[:4096]:
            raise ValueError("This text file contains binary data")
        if head.startswith((b"MZ", b"\x7fELF", b"PK\x03\x04")):
            raise ValueError("This isn't a text file")


def delete_document(persona: dict, collection, filename: str) -> int:
    """Remove an uploaded file and every memory made from it; returns how
    many memories were removed. FileNotFoundError if there's no such upload."""
    name = _safe_filename(filename)
    with _lock:
        fresh = load_persona(persona["id"], any_owner=True) or persona
        uploads = fresh.get("uploads", [])
        if not any(u["filename"] == name for u in uploads):
            raise FileNotFoundError(name)
        removed = len(collection.get(where={"source_file": name}, include=[])["ids"])
        collection.delete(where={"source_file": name})
        (CUSTOM_DIR / persona["id"] / "uploads" / name).unlink(missing_ok=True)
        fresh["uploads"] = [u for u in uploads if u["filename"] != name]
        save_persona(fresh)
    return removed


# Voice samples for cloning (XTTS-v2 needs ~6 s; longer adds little)
VOICE_MIN_SECONDS, VOICE_MAX_SECONDS = 6, 60
VOICE_CONSENT_STATEMENT = (
    "This person, or their estate, agreed to their voice being used for this model, and to the "
    "recording being sent to Fish Audio (a cloud service) to make a private voice from it. "
    "Voice is biometric data: removing the voice deletes it here and at Fish Audio."
)


def voice_path(persona: dict) -> Path:
    return CUSTOM_DIR / persona["id"] / "voice" / "reference.wav"


def check_voice_sample(data: bytes) -> float:
    """Seconds of speech in a reference recording.

    Raises ValueError unless it is a readable PCM .wav of 6-60 seconds.
    """
    import io
    import wave

    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError(f"Recording is larger than {MAX_UPLOAD_BYTES // (1024 * 1024)} MB")
    try:
        with wave.open(io.BytesIO(data)) as wav:
            seconds = wav.getnframes() / float(wav.getframerate())
    except (wave.Error, EOFError, ZeroDivisionError):
        raise ValueError("Please upload a standard (PCM) .wav recording")
    if not VOICE_MIN_SECONDS <= seconds <= VOICE_MAX_SECONDS:
        raise ValueError(f"The recording is {seconds:.0f} s; use {VOICE_MIN_SECONDS}-{VOICE_MAX_SECONDS} s of clear speech")
    return seconds


def save_voice_sample(persona: dict, data: bytes, fish_voice_id: str, statement: str = "", language: str = "en") -> dict:
    """Store a consented recording and its private Fish voice id (custom models only).

    The .wav is kept locally so the voice can be remade, e.g. with another engine.
    """
    seconds = check_voice_sample(data)
    path = voice_path(persona)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    voice = {"seconds": round(seconds, 1), "fish_voice_id": fish_voice_id,
             "consent": {"statement": statement or VOICE_CONSENT_STATEMENT, "language": language, "given_at": _now()}}
    with _lock:
        fresh = load_persona(persona["id"], any_owner=True) or persona
        fresh["voice"] = voice
        save_persona(fresh)
    return voice


def delete_voice_sample(persona: dict) -> None:
    voice_path(persona).unlink(missing_ok=True)
    with _lock:
        fresh = load_persona(persona["id"], any_owner=True) or persona
        fresh.pop("voice", None)
        save_persona(fresh)


def record_interview_answer(persona: dict, question_id: str) -> None:
    with _lock:
        fresh = load_persona(persona["id"], any_owner=True) or persona
        if question_id not in fresh.setdefault("interview_answered", []):
            fresh["interview_answered"].append(question_id)
            save_persona(fresh)


def update_persona(persona_id: str, **fields) -> dict:
    """Set fields on a persona record (read-modify-write under the lock)."""
    with _lock:
        fresh = load_persona(persona_id, any_owner=True)
        fresh.update(fields)
        save_persona(fresh)
    return fresh


def record_consent_change(persona_id: str, change: dict) -> dict:
    """Apply a consent change (services/consent.py) and add it to the model's consent history."""
    with _lock:
        fresh = load_persona(persona_id, any_owner=True)
        fresh.update(change)
        fresh.setdefault("consent_history", []).append({"at": _now(), **change})
        save_persona(fresh)
    return fresh


def record_followup(persona: dict, question_id: str) -> None:
    with _lock:
        fresh = load_persona(persona["id"], any_owner=True) or persona
        if question_id not in fresh.setdefault("followups_answered", []):
            fresh["followups_answered"].append(question_id)
            save_persona(fresh)


def build_persona(collection, persona: dict) -> dict:
    """Mark a custom persona ready after the data sufficiency check.

    Memories are embedded as they are added, so "building" is the gate that
    makes the model available for chat, not a long training job.
    """
    count = collection.count()
    if count < MIN_MEMORIES_TO_BUILD:
        raise ValueError(
            f"Only {count} memories so far; at least {MIN_MEMORIES_TO_BUILD} are needed. "
            "Upload more documents or answer more interview questions."
        )
    with _lock:
        fresh = load_persona(persona["id"], any_owner=True) or persona
        fresh.update(status="ready", built_at=_now())
        save_persona(fresh)
    return fresh


def delete_custom_persona(client, persona: dict) -> None:
    """Permanently delete a custom persona: vectors, uploads, record, and its
    questions and answers in the Q&A log (they quote its private memories)."""
    from services import qa_log

    if persona.get("kind") != "custom":
        raise PermissionError("Only custom models can be deleted")
    try:
        client.delete_collection(persona["collection"])
    except Exception:
        pass  # never had any memories
    qa_log.purge(persona["id"])
    from services.insights import purge_feedback
    purge_feedback(persona["id"])
    folder = (CUSTOM_DIR / persona["id"]).resolve()
    if folder.parent == CUSTOM_DIR.resolve() and folder.exists():
        shutil.rmtree(folder)
